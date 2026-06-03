"""
Cron: Rex Airlines scraper + email
Runs one subprocess per route, each writing its own .xlsx file.
If ALL rows in a route have no price (N/A) the route is marked FAILED.
Emails all per-route files with a full breakdown on completion.
"""

import os
import sys
import time
import smtplib
import subprocess
import threading
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from pathlib import Path

# ── Configuration ───────────────────────────────────────

OUTPUT_DIR = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)

EMAIL_FROM     = os.getenv("EMAIL_FROM", "ahteshamsalamat@gmail.com")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD", "")
EMAIL_TO       = os.getenv("EMAIL_TO", "ahteshamansari@bizprospex.com")

SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 587

MAX_RETRIES   = 3
RETRY_DELAY_S = 60
RETRY_ERRORS  = ("Connection aborted.", "RemoteDisconnected")

AIRPORT_NAMES = {
    "PER": "Perth", "ALH": "Albany", "EPR": "Esperance",
    "CVQ": "Carnarvon", "MJK": "Monkey Mia",
}

# ── Route definitions ──────────────────────────────────

ROUTES = [
    {"orig": "PER", "dest": "ALH"},
    {"orig": "ALH", "dest": "PER"},
    {"orig": "PER", "dest": "EPR"},
    {"orig": "EPR", "dest": "PER"},
    {"orig": "PER", "dest": "CVQ"},
    {"orig": "CVQ", "dest": "PER"},
    {"orig": "CVQ", "dest": "MJK"},
    {"orig": "MJK", "dest": "CVQ"},
]


def route_label(orig: str, dest: str) -> str:
    return (
        f"{orig} -> {dest} "
        f"({AIRPORT_NAMES.get(orig, orig)} -> {AIRPORT_NAMES.get(dest, dest)})"
    )


# ── Helpers ─────────────────────────────────────────────

def _au_now() -> datetime:
    """Current time in Australian (Perth) timezone — matches the cron TZ env.
    Falls back to UTC+8 manually if zoneinfo / tzdata aren't available."""
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo("Australia/Perth"))
    except Exception:
        from datetime import timezone, timedelta
        return datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=8)))


def log(msg: str) -> None:
    ts = _au_now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{ts}] {msg}", flush=True)


def format_duration(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    if h > 0:
        return f"{h}h {m}m {s}s"
    elif m > 0:
        return f"{m}m {s}s"
    return f"{s}s"


def stream_process(cmd: list, env: dict, timeout: float) -> tuple[int, str]:
    output_lines: list[str] = []
    proc = subprocess.Popen(
        cmd, env=env,
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
        text=True, bufsize=1,
    )

    def _reader():
        for line in proc.stdout:
            print(line, end="", flush=True)
            output_lines.append(line)

    t = threading.Thread(target=_reader, daemon=True)
    t.start()
    try:
        proc.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        proc.kill()
        t.join(timeout=5)
        raise
    t.join(timeout=5)
    return proc.returncode, "".join(output_lines)


def _is_no_price(val) -> bool:
    if val is None:
        return True
    s = str(val).strip()
    return s in ("", "N/A", "No Data")


def analyse_route_file(xlsx_path: Path) -> dict:
    """
    Open the per-route xlsx, count total rows and no-price rows.
    Returns total, no_price, and all_no_price flag.
    """
    if not xlsx_path.exists():
        return {"total": 0, "no_price": 0, "all_no_price": True}

    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = list(next(rows_iter, []))
        try:
            price_col = headers.index("Fare Price")
        except ValueError:
            wb.close()
            return {"total": 0, "no_price": 0, "all_no_price": True}

        total    = 0
        no_price = 0
        for row in rows_iter:
            total += 1
            if _is_no_price(row[price_col] if price_col < len(row) else None):
                no_price += 1

        wb.close()
        return {
            "total":       total,
            "no_price":    no_price,
            "all_no_price": total == 0 or no_price == total,
        }
    except Exception as e:
        log(f"  Could not read {xlsx_path.name}: {e}")
        return {"total": 0, "no_price": 0, "all_no_price": True}


# ── Per-route runner ─────────────────────────────────────

def run_route(route: dict, stamp: str) -> dict:
    orig  = route["orig"]
    dest  = route["dest"]
    label = route_label(orig, dest)
    fname = f"Rex_{orig}_{dest}_{stamp}.xlsx"
    fpath = OUTPUT_DIR / fname
    cmd   = [
        "python", "rex_brightdata.py",
        "--routes", f"{orig}-{dest}",
        "--output", str(fpath),
    ]

    log(f"{'=' * 55}")
    log(f"Starting route: {label}")
    log(f"   Output : {fname}")
    log(f"   Command: {' '.join(cmd)}")
    log("")

    start     = time.time()
    success   = False
    exit_code = -1
    duration  = "0s"

    for attempt in range(1, MAX_RETRIES + 1):
        if attempt > 1:
            log(f"Retry {attempt}/{MAX_RETRIES} for {label} -- waiting {RETRY_DELAY_S}s...")
            time.sleep(RETRY_DELAY_S)

        try:
            env = os.environ.copy()
            env["PYTHONUNBUFFERED"] = "1"
            env["TZ"] = "Australia/Perth"

            returncode, output = stream_process(cmd, env, timeout=14400)
            elapsed   = time.time() - start
            duration  = format_duration(elapsed)
            exit_code = returncode
            success   = returncode == 0

            if success:
                log(f"{label} subprocess completed in {duration}")
                break

            log(f"{label} subprocess failed (exit {returncode}) after {duration}")
            if any(err in output for err in RETRY_ERRORS):
                log("   Connection error -- will retry.")
                if attempt < MAX_RETRIES:
                    continue
            break

        except subprocess.TimeoutExpired:
            duration  = format_duration(time.time() - start)
            log(f"{label} timed out after {duration}")
            exit_code = -1
            success   = False
            break

        except Exception as e:
            duration  = format_duration(time.time() - start)
            log(f"{label} crashed: {e}")
            exit_code = -1
            success   = False
            break

    # Even if subprocess succeeded, treat as failed if all prices are N/A
    file_stats = analyse_route_file(fpath)
    if success and file_stats["all_no_price"]:
        log(f"  {label}: all rows have no price -- marking as FAILED")
        success = False

    price_ok = file_stats["total"] - file_stats["no_price"]
    pct      = (price_ok / file_stats["total"] * 100) if file_stats["total"] else 0
    log(
        f"  {label}: {file_stats['total']} rows, "
        f"{price_ok} with price ({pct:.0f}%), "
        f"{file_stats['no_price']} no price"
    )

    return {
        "orig":          orig,
        "dest":          dest,
        "label":         label,
        "file":          fpath,
        "success":       success,
        "exit_code":     exit_code,
        "duration":      duration,
        "total_rows":    file_stats["total"],
        "no_price_rows": file_stats["no_price"],
        "all_no_price":  file_stats["all_no_price"],
    }


# ── Email ────────────────────────────────────────────────

def build_email_body(route_results: list[dict], overall_success: bool) -> str:
    today = _au_now().strftime("%A, %d %B %Y")
    lines = [
        f"Flight Scraper Report -- Rex Airlines -- {today}",
        "=" * 62, "",
        f"Status   : {'Completed' if overall_success else 'FAILED (one or more routes)'}",
        f"Routes   : {len(route_results)}",
        "",
        "-" * 62,
        "Per-Route Breakdown",
        "-" * 62,
    ]

    for r in route_results:
        price_ok = r["total_rows"] - r["no_price_rows"]
        pct      = (price_ok / r["total_rows"] * 100) if r["total_rows"] else 0
        r_status = (
            "FAILED - No Data" if r["all_no_price"] else
            "FAILED"          if not r["success"]   else
            "PARTIAL"         if r["no_price_rows"] > 0 else
            "OK"
        )
        lines += [
            f"  [{r_status}] {r['label']}",
            f"       File        : {r['file'].name}",
            f"       Duration    : {r['duration']}",
            f"       Total rows  : {r['total_rows']}",
            f"       With price  : {price_ok}  ({pct:.0f}%)",
            f"       No price    : {r['no_price_rows']}",
            "",
        ]

    files_exist = [r for r in route_results if r["file"].exists()]
    if files_exist:
        lines += [
            "-" * 62,
            f"Attached files ({len(files_exist)}):",
        ]
        for r in files_exist:
            size_kb = r["file"].stat().st_size / 1024
            lines.append(f"  - {r['file'].name}  ({size_kb:.1f} KB)")

    lines.append("")
    return "\n".join(lines)


def send_email(route_results: list[dict], overall_success: bool) -> None:
    if not EMAIL_PASSWORD:
        log("EMAIL_PASSWORD not set -- skipping email.")
        return

    today   = _au_now().strftime("%Y-%m-%d")
    status  = "OK" if overall_success else "FAILED"
    subject = f"Rex Airlines Scraper -- {today} -- {status}"
    body    = build_email_body(route_results, overall_success)

    msg = MIMEMultipart()
    msg["From"]    = EMAIL_FROM
    msg["To"]      = EMAIL_TO
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain", "utf-8"))

    for r in route_results:
        filepath = r["file"]
        if not filepath.exists():
            log(f"File missing, skipping attachment: {filepath.name}")
            continue
        try:
            with open(filepath, "rb") as f:
                part = MIMEBase("application", "octet-stream")
                part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f'attachment; filename="{filepath.name}"')
            msg.attach(part)
        except Exception as e:
            log(f"Could not attach {filepath.name}: {e}")

    attached = sum(1 for r in route_results if r["file"].exists())
    log(f"Sending email to {EMAIL_TO} ({attached} attachments)...")
    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
            server.ehlo()
            server.starttls()
            server.ehlo()
            server.login(EMAIL_FROM, EMAIL_PASSWORD)
            server.send_message(msg)
        log("Email sent successfully!")
    except Exception as e:
        log(f"Email failed: {e}")


# ── Main ─────────────────────────────────────────────────

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Rex Airlines cron: scrape per-route + email")
    parser.add_argument("--dry-run", action="store_true",
                        help="Skip scrapers, email existing output files")
    args = parser.parse_args()

    log("=" * 55)
    log("Rex Airlines Scraper Cron")
    log(f"   Date  : {_au_now().strftime('%A, %d %B %Y %H:%M %Z')}")
    log(f"   Mode  : {'DRY RUN' if args.dry_run else 'FULL RUN'}")
    log(f"   Routes: {len(ROUTES)}")
    log("=" * 55)
    log("")

    stamp         = datetime.now().strftime("%d-%m-%Y_%H-%M")
    route_results = []
    any_failed    = False

    for i, route in enumerate(ROUTES, 1):
        log(f"[Route {i}/{len(ROUTES)}] {route_label(route['orig'], route['dest'])}")

        if args.dry_run:
            fname = f"Rex_{route['orig']}_{route['dest']}_{stamp}.xlsx"
            fpath = OUTPUT_DIR / fname
            stats = analyse_route_file(fpath) if fpath.exists() else {"total": 0, "no_price": 0, "all_no_price": True}
            result = {
                "orig":          route["orig"],
                "dest":          route["dest"],
                "label":         route_label(route["orig"], route["dest"]),
                "file":          fpath,
                "success":       not stats["all_no_price"],
                "exit_code":     0,
                "duration":      "dry-run",
                "total_rows":    stats["total"],
                "no_price_rows": stats["no_price"],
                "all_no_price":  stats["all_no_price"],
            }
        else:
            result = run_route(route, stamp)

        route_results.append(result)
        if not result["success"]:
            any_failed = True

        log("")

    overall_success = not any_failed

    log("=" * 55)
    log("Route Summary")
    log("=" * 55)
    for r in route_results:
        tag = "OK" if r["success"] else "FAILED"
        log(f"  [{tag}] {r['label']}  ({r['duration']})")
    log("")

    send_email(route_results, overall_success)

    log("")
    log("=" * 55)
    log(f"Done -- {'Success' if overall_success else 'FAILED (one or more routes)'}")
    log("=" * 55)

    if not overall_success:
        sys.exit(1)


if __name__ == "__main__":
    main()
