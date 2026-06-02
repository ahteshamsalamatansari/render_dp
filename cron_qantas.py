"""
Cron: Qantas scraper + email
Runs qantas_playwright2ndJune.py (all 8 routes, built-in first-pass + retry).
After completion, finds per-route xlsx files, counts no-price rows per route,
and emails all files with a full per-route breakdown.
"""

import os
import sys
import time
import smtplib
import subprocess
import threading
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from pathlib import Path

# ── Configuration ───────────────────────────────────────

OUTPUT_DIR = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)

EMAIL_FROM     = os.getenv("EMAIL_FROM", "ahteshamsalamat@gmail.com")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD", "")
EMAIL_TO       = os.getenv("EMAIL_TO", "ahteshamansari@bizprospex.com")

SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 587

SCRAPER_CMD = ["python", "qantas_playwright2ndJune.py"]

AIRPORT_NAMES = {
    "BME": "Broome", "KNX": "Kununurra", "DRW": "Darwin",
    "PER": "Perth",  "GET": "Geraldton",
}

# Must match ROUTES order in qantas_playwright2ndJune.py
ROUTES = [
    {"orig": "BME", "dest": "KNX"},
    {"orig": "BME", "dest": "DRW"},
    {"orig": "DRW", "dest": "KNX"},
    {"orig": "KNX", "dest": "BME"},
    {"orig": "PER", "dest": "GET"},
    {"orig": "GET", "dest": "PER"},
    {"orig": "DRW", "dest": "BME"},
    {"orig": "KNX", "dest": "DRW"},
]

NO_PRICE_CLASSES = {"No Direct Flight", "NO FLIGHTS", "SOLD OUT", "NO DATA"}

# ── Helpers ─────────────────────────────────────────────

def log(msg: str) -> None:
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{ts}] {msg}", flush=True)


def format_duration(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    if h > 0:
        return f"{h}h {m}m {s}s"
    elif m > 0:
        return f"{m}m {s}s"
    return f"{s}s"


def route_label(orig: str, dest: str) -> str:
    return (
        f"{orig} -> {dest} "
        f"({AIRPORT_NAMES.get(orig, orig)} -> {AIRPORT_NAMES.get(dest, dest)})"
    )


def stream_process(cmd: list, env: dict, timeout: float) -> tuple[int, str]:
    output_lines: list[str] = []
    proc = subprocess.Popen(
        cmd, env=env,
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
        text=True, bufsize=1,
    )

    def _reader():
        for line in proc.stdout:
            print(line, end="", flush=True)
            output_lines.append(line)

    t = threading.Thread(target=_reader, daemon=True)
    t.start()
    try:
        proc.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        proc.kill()
        t.join(timeout=5)
        raise
    t.join(timeout=5)
    return proc.returncode, "".join(output_lines)


def collect_output_files_since(since_ts: float) -> list[Path]:
    files = []
    if not OUTPUT_DIR.exists():
        return files
    for item in OUTPUT_DIR.rglob("*"):
        if item.is_file() and item.suffix.lower() in (".csv", ".xlsx"):
            if item.stat().st_mtime >= since_ts:
                files.append(item)
    files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    return files


# ── Per-route file analysis ──────────────────────────────

def find_route_file(orig: str, dest: str, ext: str) -> "Path | None":
    """Find the most recent Qantas_{orig}-{dest}_*.{ext} in output/."""
    candidates = sorted(
        OUTPUT_DIR.glob(f"Qantas_{orig}-{dest}_*.{ext}"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None


def _is_no_price(fare_price, fare_class: str) -> bool:
    if fare_price is None:
        return True
    s = str(fare_price).strip()
    if s in ("", "N/A", "No Data", "None", "nan"):
        return True
    if fare_class in NO_PRICE_CLASSES:
        return True
    try:
        return float(s) == 0.0
    except ValueError:
        return False


def analyse_route_xlsx(xlsx_path: Path) -> dict:
    """Read route xlsx, count total rows and no-price rows."""
    if not xlsx_path or not xlsx_path.exists():
        return {"total": 0, "no_price": 0, "all_no_price": True}
    try:
        from openpyxl import load_workbook
        wb  = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws  = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers   = list(next(rows_iter, []))
        try:
            price_col = headers.index("Fare Price")
            class_col = headers.index("Fare Class")
        except ValueError:
            wb.close()
            return {"total": 0, "no_price": 0, "all_no_price": True}

        total    = 0
        no_price = 0
        for row in rows_iter:
            total += 1
            fp = row[price_col] if price_col < len(row) else None
            fc = str(row[class_col]) if class_col < len(row) and row[class_col] else ""
            if _is_no_price(fp, fc):
                no_price += 1

        wb.close()
        return {
            "total":        total,
            "no_price":     no_price,
            "all_no_price": total == 0 or no_price == total,
        }
    except Exception as e:
        log(f"  Could not read {xlsx_path.name}: {e}")
        return {"total": 0, "no_price": 0, "all_no_price": True}


def build_route_stats() -> list[dict]:
    """Match each route to its output files and compute price coverage stats."""
    stats = []
    for r in ROUTES:
        orig, dest = r["orig"], r["dest"]
        xlsx       = find_route_file(orig, dest, "xlsx")
        csv_file   = find_route_file(orig, dest, "csv")
        fs         = analyse_route_xlsx(xlsx)

        price_ok = fs["total"] - fs["no_price"]
        pct      = (price_ok / fs["total"] * 100) if fs["total"] else 0
        r_status = (
            "FAILED - No File" if not xlsx else
            "FAILED - No Data" if fs["all_no_price"] else
            "PARTIAL"          if fs["no_price"] > 0 else
            "OK"
        )
        log(
            f"   {route_label(orig, dest)}: {fs['total']} rows, "
            f"{price_ok} with price ({pct:.0f}%), {fs['no_price']} no price  [{r_status}]"
        )
        stats.append({
            "orig":          orig,
            "dest":          dest,
            "label":         route_label(orig, dest),
            "xlsx":          xlsx,
            "csv":           csv_file if csv_file and csv_file.exists() else None,
            "total_rows":    fs["total"],
            "no_price_rows": fs["no_price"],
            "all_no_price":  fs["all_no_price"],
            "status":        r_status,
        })
    return stats


# ── Email ────────────────────────────────────────────────

def build_email_body(scraper_success: bool, duration: str, route_stats: list[dict]) -> str:
    today   = datetime.now().strftime("%A, %d %B %Y")
    any_bad = any(s["status"] != "OK" for s in route_stats)
    overall = "FAILED" if (not scraper_success or any_bad) else "Completed"
    lines = [
        f"Flight Scraper Report -- Qantas -- {today}",
        "=" * 62, "",
        f"Status   : {overall}",
        f"Duration : {duration}",
        f"Routes   : {len(route_stats)}",
        "",
        "-" * 62,
        "Per-Route Breakdown",
        "-" * 62,
    ]

    for s in route_stats:
        price_ok = s["total_rows"] - s["no_price_rows"]
        pct      = (price_ok / s["total_rows"] * 100) if s["total_rows"] else 0
        lines += [
            f"  [{s['status']}] {s['label']}",
            f"       File        : {s['xlsx'].name if s['xlsx'] else 'NOT FOUND'}",
            f"       Total rows  : {s['total_rows']}",
            f"       With price  : {price_ok}  ({pct:.0f}%)",
            f"       No price    : {s['no_price_rows']}",
            "",
        ]

    all_files = [
        f for s in route_stats
        for f in [s["xlsx"], s["csv"]]
        if f and f.exists()
    ]
    if all_files:
        lines += [
            "-" * 62,
            f"Attached files ({len(all_files)}):",
        ]
        for f in all_files:
            size_kb = f.stat().st_size / 1024
            lines.append(f"  - {f.name}  ({size_kb:.1f} KB)")
    else:
        lines.append("  No output files were generated.")

    lines.append("")
    return "\n".join(lines)


def send_email(scraper_success: bool, duration: str, route_stats: list[dict]) -> None:
    if not EMAIL_PASSWORD:
        log("EMAIL_PASSWORD not set -- skipping email.")
        return

    any_bad = any(s["status"] != "OK" for s in route_stats)
    today   = datetime.now().strftime("%Y-%m-%d")
    status  = "FAILED" if (not scraper_success or any_bad) else "OK"
    subject = f"Qantas Scraper -- {today} -- {status}"
    body    = build_email_body(scraper_success, duration, route_stats)

    msg = MIMEMultipart()
    msg["From"]    = EMAIL_FROM
    msg["To"]      = EMAIL_TO
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain", "utf-8"))

    for s in route_stats:
        for filepath in [s["xlsx"], s["csv"]]:
            if not filepath or not filepath.exists():
                continue
            try:
                with open(filepath, "rb") as f:
                    part = MIMEBase("application", "octet-stream")
                    part.set_payload(f.read())
                encoders.encode_base64(part)
                part.add_header("Content-Disposition", f'attachment; filename="{filepath.name}"')
                msg.attach(part)
            except Exception as e:
                log(f"Could not attach {filepath.name}: {e}")

    total_attached = sum(
        1 for s in route_stats
        for f in [s["xlsx"], s["csv"]]
        if f and f.exists()
    )
    log(f"Sending email to {EMAIL_TO} ({total_attached} attachments)...")
    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
            server.ehlo()
            server.starttls()
            server.ehlo()
            server.login(EMAIL_FROM, EMAIL_PASSWORD)
            server.send_message(msg)
        log("Email sent successfully!")
    except Exception as e:
        log(f"Email failed: {e}")


# ── Runner ───────────────────────────────────────────────

def run_scraper() -> tuple[bool, str]:
    """Run qantas_playwright2ndJune.py — handles all 8 routes with built-in retry."""
    log(f"{'=' * 55}")
    log(f"Starting Qantas scraper...")
    log(f"   Command : {' '.join(SCRAPER_CMD)}")
    log(f"   Routes  : {len(ROUTES)} (scraper handles first-pass + retry internally)")
    log("")

    start = time.time()
    try:
        env = os.environ.copy()
        env["PYTHONUNBUFFERED"] = "1"
        env["TZ"] = "Australia/Perth"

        returncode, _ = stream_process(SCRAPER_CMD, env, timeout=21600)  # 6h hard cap
        duration = format_duration(time.time() - start)
        success  = returncode == 0

        if success:
            log(f"Scraper completed in {duration}")
        else:
            log(f"Scraper finished with exit code {returncode} after {duration}")
        return success, duration

    except subprocess.TimeoutExpired:
        duration = format_duration(time.time() - start)
        log(f"Scraper hard-timeout after {duration}")
        return False, duration

    except Exception as e:
        duration = format_duration(time.time() - start)
        log(f"Scraper crashed: {e}")
        return False, duration


# ── Main ─────────────────────────────────────────────────

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Qantas cron: scrape + email")
    parser.add_argument("--dry-run", action="store_true",
                        help="Skip scraper, analyse existing files and email")
    args = parser.parse_args()

    log("=" * 55)
    log("Qantas Scraper Cron")
    log(f"   Date  : {datetime.now().strftime('%A, %d %B %Y %H:%M %Z')}")
    log(f"   Mode  : {'DRY RUN' if args.dry_run else 'FULL RUN'}")
    log("=" * 55)
    log("")

    if args.dry_run:
        log("Dry run -- skipping scraper.")
        scraper_success = True
        duration        = "dry-run"
    else:
        scraper_success, duration = run_scraper()

    log("\nAnalysing per-route output files...")
    route_stats = build_route_stats()

    log(f"\nRoute Summary:")
    log("=" * 55)
    for s in route_stats:
        log(f"  [{s['status']}] {s['label']}")
    log("")

    send_email(scraper_success, duration, route_stats)

    log("")
    log("=" * 55)
    any_bad = any(s["status"] != "OK" for s in route_stats)
    log(f"Done -- {'FAILED' if (not scraper_success or any_bad) else 'Success'}")
    log("=" * 55)

    if not scraper_success or any_bad:
        sys.exit(1)


if __name__ == "__main__":
    main()
